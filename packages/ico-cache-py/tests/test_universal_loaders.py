import os
import pytest
from ico_cache.loaders.base import BaseLoader, Chunk
from ico_cache.loaders.structured_loader import StructuredLoader
from ico_cache.loaders.code_loader import CodeLoader
from ico_cache.loaders.txt_loader import TXTLoader
from ico_cache.loaders.html_loader import HTMLLoader
from ico_cache.loaders.auto_loader import AutoLoader, ingest
from examples.universal_schema import universal_schema

from fixtures_gen import generate_corpus


@pytest.fixture(scope="session")
def corpus(tmp_path_factory):
    """Generate the entire synthetic corpus once per session into a tmp dir."""
    root = tmp_path_factory.mktemp("corpus")
    return generate_corpus(str(root))


def test_structured_loader_csv(corpus):
    loader = StructuredLoader()
    csv_path = os.path.join(corpus, "data/products.csv")
    chunks = loader.load(csv_path)
    assert len(chunks) == 3
    assert all(isinstance(c, Chunk) for c in chunks)
    assert chunks[0].metadata["name"] == "Sony WH-1000XM5"
    assert "P101" in chunks[0].text


def test_structured_loader_json():
    loader = StructuredLoader()
    raw_json = '[{"user_id": 1, "action": "login"}, {"user_id": 2, "action": "logout"}]'
    chunks = loader.load(raw_json)
    assert len(chunks) == 2
    assert chunks[0].metadata["user_id"] == 1
    assert "action: login" in chunks[0].text


def test_code_loader_ast(corpus):
    loader = CodeLoader(language="python")
    code_path = os.path.join(corpus, "data/sample_code.py")
    chunks = loader.load(code_path)
    assert len(chunks) >= 2
    func_chunk = next(c for c in chunks if c.metadata.get("name") == "calculate_discount")
    assert func_chunk.metadata["type"] == "function"
    assert "def calculate_discount" in func_chunk.text

    class_chunk = next(c for c in chunks if c.metadata.get("name") == "InventoryManager")
    assert class_chunk.metadata["type"] == "class"
    assert "class InventoryManager" in class_chunk.text


def test_txt_loader(corpus):
    loader = TXTLoader()
    txt_path = os.path.join(corpus, "data/sample.txt")
    chunks = loader.load(txt_path)
    assert len(chunks) == 3
    assert "Universal Document Ingestion" in chunks[0].text


def test_ocr_disabled_returns_no_ocr():
    """
    Verify configure_ocr(enabled=False) disables content-based OCR fallback,
    independent of the file type.
    """
    import tempfile
    from ico_cache.loaders.ocr import configure_ocr, ocr_any

    with tempfile.TemporaryDirectory() as td:
        try:
            from PIL import Image, ImageDraw
        except ImportError:
            pytest.skip("Pillow not installed")
        png = os.path.join(td, "scan.png")
        image = Image.new("RGB", (400, 80), "white")
        ImageDraw.Draw(image).text((10, 30), "Invoice", fill="black")
        image.save(png)

        try:
            configure_ocr(enabled=False)
            assert ocr_any(png) == []
        finally:
            configure_ocr(enabled=True)


def _tesseract_available() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def test_odf_formats_universal(tmp_path):
    """
    Verify OpenDocument text (.odt) and spreadsheet (.ods) ingests through the
    universal AutoLoader: body paragraphs, headings and table rows all become chunks.
    """
    odf = pytest.importorskip("odf")
    from odf.opendocument import OpenDocumentText, OpenDocumentSpreadsheet
    from odf.text import P, H
    from odf.table import Table, TableRow, TableCell

    auto = AutoLoader(schema=universal_schema)

    # .odt with a paragraph, a heading and a table
    doc = OpenDocumentText()
    doc.text.addElement(H(outlinelevel=1, text="Quarterly Report"))
    doc.text.addElement(P(text="Revenue grew across every region in Q3."))
    table = Table(name="Q3")
    header = TableRow()
    for value in ("region", "revenue"):
        cell = TableCell()
        cell.addElement(P(text=value))
        header.addElement(cell)
    table.addElement(header)
    for region, revenue in (("North", "1.2M"), ("South", "0.9M")):
        row = TableRow()
        for value in (region, revenue):
            cell = TableCell()
            cell.addElement(P(text=value))
            row.addElement(cell)
        table.addElement(row)
    doc.text.addElement(table)

    odt_path = os.path.join(tmp_path, "report.odt")
    doc.save(odt_path)
    odt_chunks = auto.load(odt_path)
    odt_text = "\n".join(c.text for c in odt_chunks)
    assert "Quarterly Report" in odt_text
    assert "Revenue grew across every region" in odt_text
    assert "region: North" in odt_text
    assert any(c.loader_type == "odf_table" for c in odt_chunks)

    # .ods spreadsheet
    sheet_doc = OpenDocumentSpreadsheet()
    sheet = Table(name="Sheet1")
    for row_values in (("item", "price"), ("widget", "9.99"), ("gadget", "19.99")):
        row = TableRow()
        for value in row_values:
            cell = TableCell()
            cell.addElement(P(text=value))
            row.addElement(cell)
        sheet.addElement(row)
    sheet_doc.spreadsheet.addElement(sheet)

    ods_path = os.path.join(tmp_path, "prices.ods")
    sheet_doc.save(ods_path)
    ods_chunks = auto.load(ods_path)
    ods_text = "\n".join(c.text for c in ods_chunks)
    assert "item: widget" in ods_text
    assert "price: 19.99" in ods_text


def test_office_open_xml_formats(tmp_path):
    """
    Verify .docx, .xlsx and .pptx ingest through the universal AutoLoader.
    Skips formats whose optional library is not installed.
    """
    auto = AutoLoader(schema=universal_schema)
    produced = 0

    try:
        import docx

        document = docx.Document()
        document.add_paragraph("Quarterly financial outlook for the company.")
        docx_path = os.path.join(tmp_path, "report.docx")
        document.save(docx_path)
        chunks = auto.load(docx_path)
        assert any("Quarterly financial outlook" in c.text for c in chunks)
        produced += 1
    except ImportError:
        pass

    try:
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.append(["metric", "value"])
        worksheet.append(["operating_margin", "28.4%"])
        xlsx_path = os.path.join(tmp_path, "metrics.xlsx")
        workbook.save(xlsx_path)
        chunks = auto.load(xlsx_path)
        assert any("operating_margin" in c.text and "28.4%" in c.text for c in chunks)
        produced += 1
    except ImportError:
        pass

    try:
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "Capital allocation summary"
        pptx_path = os.path.join(tmp_path, "deck.pptx")
        presentation.save(pptx_path)
        chunks = auto.load(pptx_path)
        assert any("Capital allocation summary" in c.text for c in chunks)
        produced += 1
    except ImportError:
        pass

    if produced == 0:
        pytest.skip("No office libraries (python-docx/openpyxl/python-pptx) installed")


def test_ocr_content_sniffing_ignores_extension(tmp_path):
    """
    Verify OCR is content-based: a PNG renamed to an unknown extension is still
    OCR'd, and an image with a misleading extension is handled via magic bytes.
    """
    pytest.importorskip("PIL")
    from PIL import Image, ImageDraw

    if not _tesseract_available():
        pytest.skip("tesseract binary not available")

    image = Image.new("RGB", (900, 160), "white")
    ImageDraw.Draw(image).text((20, 60), "Invoice Total Amount Due", fill="black")
    disguised = os.path.join(tmp_path, "mystery.dat")
    image.save(disguised, format="PNG")

    chunks = AutoLoader().load(disguised)
    assert chunks, "content-based OCR should read a PNG regardless of extension"
    text = " ".join(c.text for c in chunks)
    assert "Invoice" in text
    assert chunks[0].loader_type == "image_ocr"
    assert chunks[0].metadata.get("extraction_method") == "ocr"


def test_image_loader_dispatches_by_extension(tmp_path):
    """Verify common image extensions route to the OCR ImageLoader."""
    pytest.importorskip("PIL")
    from PIL import Image, ImageDraw
    from ico_cache.loaders.image_loader import ImageLoader

    if not _tesseract_available():
        pytest.skip("tesseract binary not available")

    image = Image.new("RGB", (900, 160), "white")
    ImageDraw.Draw(image).text((20, 60), "Scanned Statement Balance", fill="black")
    for ext in (".png", ".jpg", ".tiff"):
        path = os.path.join(tmp_path, f"scan{ext}")
        image.save(path)
        chunks = ImageLoader().load(path)
        assert chunks, f"expected OCR text for {ext}"
        assert "Statement" in chunks[0].text


def test_unknown_extension_text_fallback(tmp_path):
    """Unknown extensions with text content are read as plain text; binary is ignored."""
    text_path = os.path.join(tmp_path, "notes.custom")
    with open(text_path, "w") as f:
        f.write("Universal fallback reads unknown text extensions.\n\nSecond paragraph here.")
    chunks = AutoLoader().load(text_path)
    assert len(chunks) == 2
    assert "Universal fallback" in chunks[0].text

    binary_path = os.path.join(tmp_path, "blob.custom")
    with open(binary_path, "wb") as f:
        f.write(bytes(range(256)) * 8)
    assert AutoLoader().load(binary_path) == []


def test_auto_loader_dispatch(corpus):
    txt_path = os.path.join(corpus, "data/sample.txt")
    csv_path = os.path.join(corpus, "data/products.csv")
    chunks = ingest(path=txt_path)
    assert len(chunks) == 3
    csv_chunks = ingest(path=csv_path)
    assert len(csv_chunks) == 3


def test_corpus_edge_cases_no_unhandled_exceptions(corpus):
    """
    Verify that across all corpus edge cases (empty, tiny, non-UTF-8, malformed, large),
    no loader raises an unhandled exception. Malformed input logs/skips and never crashes.
    """
    loaders = {
        "text": TXTLoader(),
        "structured": StructuredLoader(),
        "code": CodeLoader(),
    }
    auto_loader = AutoLoader()

    for category, loader in loaders.items():
        cat_dir = os.path.join(corpus, category)
        assert os.path.exists(cat_dir), f"Corpus directory missing: {cat_dir}"
        files = os.listdir(cat_dir)
        assert len(files) >= 5, f"Expected at least 5 files in {cat_dir}, found {len(files)}"

        for f in files:
            file_path = os.path.join(cat_dir, f)
            # Must not raise unhandled exception
            chunks = loader.load(file_path)
            assert isinstance(chunks, list)
            auto_chunks = auto_loader.load(file_path)
            assert isinstance(auto_chunks, list)


def test_corpus_sane_chunk_bounds_per_file_size(corpus):
    """
    Verify that chunk counts fall within a sane min/max bound per file size.
    """
    categories = {
        "text": TXTLoader(),
        "structured": StructuredLoader(),
        "code": CodeLoader(),
    }

    for cat, loader in categories.items():
        cat_dir = os.path.join(corpus, cat)

        # Empty file -> 0 chunks
        empty_files = [f for f in os.listdir(cat_dir) if f.startswith("empty")]
        for ef in empty_files:
            chunks = loader.load(os.path.join(cat_dir, ef))
            assert len(chunks) == 0, f"Expected 0 chunks for empty file {ef}, got {len(chunks)}"

        # Tiny file -> exactly 1 chunk
        tiny_files = [f for f in os.listdir(cat_dir) if f.startswith("tiny")]
        for tf in tiny_files:
            chunks = loader.load(os.path.join(cat_dir, tf))
            assert len(chunks) == 1, f"Expected 1 chunk for tiny file {tf}, got {len(chunks)}"

        # Standard file -> bounded between 2 and 20 chunks
        std_files = [f for f in os.listdir(cat_dir) if f.startswith("standard")]
        for sf in std_files:
            chunks = loader.load(os.path.join(cat_dir, sf))
            assert 2 <= len(chunks) <= 20, f"Expected 2-20 chunks for standard file {sf}, got {len(chunks)}"

        # Large file (>10MB) -> bounded between 1,000 and 150,000 chunks
        large_files = [f for f in os.listdir(cat_dir) if f.startswith("large.")]
        for lf in large_files:
            chunks = loader.load(os.path.join(cat_dir, lf))
            assert 1000 <= len(chunks) <= 150000, f"Expected 1000-150000 chunks for large file {lf}, got {len(chunks)}"


def test_corpus_schema_metadata_matching(corpus):
    """
    Verify that every chunk produced has non-null source metadata matching the configured schema.
    """
    loaders = [
        TXTLoader(schema=universal_schema),
        StructuredLoader(schema=universal_schema),
        CodeLoader(schema=universal_schema),
    ]
    test_files = [
        os.path.join(corpus, "text/standard.txt"),
        os.path.join(corpus, "structured/standard.jsonl"),
        os.path.join(corpus, "code/standard.py"),
    ]

    for loader, file_path in zip(loaders, test_files):
        chunks = loader.load(file_path, schema=universal_schema)
        assert len(chunks) > 0
        for chunk in chunks:
            assert chunk.source_file == file_path
            assert isinstance(chunk.metadata, dict)
            assert "source_file" in chunk.metadata
            assert chunk.metadata["source_file"] == file_path


def test_structural_boundary_preservation(corpus):
    """
    Verify structural boundaries:
    - StructuredLoader never splits a CSV row or JSON record across chunks.
    - CodeLoader never splits a function or class body across chunks.
    """
    # 1. Structured boundary check
    struct_loader = StructuredLoader()
    jsonl_path = os.path.join(corpus, "structured/standard.jsonl")
    jsonl_chunks = struct_loader.load(jsonl_path)
    assert len(jsonl_chunks) == 4  # Exactly 4 lines/records
    for idx, c in enumerate(jsonl_chunks):
        assert f"record_id: REC-00{idx+1}" in c.text
        assert "Row " in c.page_or_section or "Record " in c.page_or_section

    csv_path = os.path.join(corpus, "structured/tiny.csv")
    csv_chunks = struct_loader.load(csv_path)
    assert len(csv_chunks) == 1
    assert "metric: operating_margin" in csv_chunks[0].text
    assert "value: 28.4%" in csv_chunks[0].text

    # 2. CodeLoader function/class boundary check
    code_loader = CodeLoader(language="python")
    code_path = os.path.join(corpus, "code/standard.py")
    code_chunks = code_loader.load(code_path)

    # Class definition should contain all its methods intact
    order_proc = next(c for c in code_chunks if c.metadata.get("name") == "OrderProcessor")
    assert order_proc.metadata["type"] == "class"
    assert "class OrderProcessor:" in order_proc.text
    assert "def calculate_subtotal" in order_proc.text
    assert "def calculate_tax" in order_proc.text
    assert "def process_order" in order_proc.text
    assert '"status": "APPROVED"' in order_proc.text  # Full body retained, never truncated


def test_pdf_and_html_formats(corpus):
    """
    Verify TextLoader format coverage:
    - Clean PDF extracts text layer and schema metadata
    - Scanned/OCR PDF extracts text via OCR fallback, recording extraction_method='ocr'
    - Blank image PDF returns 0 chunks with explicit status 'image_only_no_text'
    - Zero-byte PDF returns 0 chunks with explicit status 'empty_file'
    - Malformed HTML strips malicious scripts and extracts clean content chunks
    Requires pymupdf (fitz) — installed as a CI dep.
    """
    pytest.importorskip("fitz")
    pytest.importorskip("PIL")
    auto_loader = AutoLoader(schema=universal_schema)
    from ico_cache.loaders.pdf_loader import PDFLoader
    import tempfile
    import fitz

    clean_pdf_path = os.path.join(corpus, "text/clean.pdf")
    scanned_pdf_path = os.path.join(corpus, "text/scanned_ocr.pdf")
    if os.path.getsize(clean_pdf_path) == 0 or os.path.getsize(scanned_pdf_path) == 0:
        pytest.skip("PDF fixtures unavailable (PyMuPDF/Pillow missing at corpus generation time)")

    # 1. Clean PDF
    pdf_chunks = auto_loader.load(clean_pdf_path)
    assert len(pdf_chunks) >= 1
    assert "Apple Inc. (AAPL)" in pdf_chunks[0].text
    assert pdf_chunks[0].metadata.get("entity") == "AAPL"
    assert pdf_chunks[0].metadata.get("extraction_method") == "text_layer"

    # 2. Scanned OCR PDF (image containing text)
    try:
        import pdf2image  # noqa: F401
    except ImportError:
        pytest.skip("pdf2image not installed (scanned PDF OCR fallback)")
    if not _tesseract_available():
        pytest.skip("tesseract binary not available (scanned PDF OCR fallback)")
    scanned_chunks = auto_loader.load(scanned_pdf_path)
    assert len(scanned_chunks) >= 1
    assert "Scanned Document" in scanned_chunks[0].text
    assert scanned_chunks[0].metadata.get("extraction_method") == "ocr"

    # 3. Blank image PDF (image with no text)
    pdf_loader = PDFLoader()
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_blank:
        tmp_blank_name = tmp_blank.name
    try:
        doc = fitz.open()
        page = doc.new_page(width=300, height=300)
        pix = fitz.Pixmap(fitz.csRGB, (0, 0, 300, 300), False)
        pix.clear_with(255)
        page.insert_image(fitz.Rect(0, 0, 300, 300), pixmap=pix)
        doc.save(tmp_blank_name)
        doc.close()

        blank_chunks = pdf_loader.load(tmp_blank_name)
        assert len(blank_chunks) == 0
        assert pdf_loader.last_status == "image_only_no_text"
    finally:
        if os.path.exists(tmp_blank_name):
            os.remove(tmp_blank_name)

    # 4. Zero-byte PDF
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_empty:
        tmp_empty_name = tmp_empty.name
    try:
        empty_chunks = pdf_loader.load(tmp_empty_name)
        assert len(empty_chunks) == 0
        assert pdf_loader.last_status == "empty_file"
    finally:
        if os.path.exists(tmp_empty_name):
            os.remove(tmp_empty_name)

    # 5. Malformed HTML (unclosed tags, script injection)
    html_path = os.path.join(corpus, "text/malformed.html")
    html_chunks = auto_loader.load(html_path)
    assert len(html_chunks) >= 1
    full_html_text = "\n".join(c.text for c in html_chunks)
    assert "Microsoft Corporation (MSFT)" in full_html_text
    assert "XSS Attack!" not in full_html_text
    assert "onerror=" not in full_html_text
    assert "Confidential copyright footer" not in full_html_text


def test_multi_language_code_loaders(corpus):
    """
    Verify tree-sitter code loading across JavaScript and Go:
    - Functions and classes are structurally bounded
    - Never split across chunks
    Requires tree-sitter-javascript and tree-sitter-go (installed as CI deps).
    """
    auto_loader = AutoLoader(schema=universal_schema)

    # JavaScript via tree-sitter
    js_path = os.path.join(corpus, "code/standard.js")
    js_chunks = auto_loader.load(js_path)
    assert len(js_chunks) == 3
    js_names = [c.metadata.get("name") for c in js_chunks]
    assert "calculateRevenue" in js_names
    assert "applyDiscount" in js_names
    assert "FinancialService" in js_names

    # Check JS class body integrity
    service_chunk = next(c for c in js_chunks if c.metadata.get("name") == "FinancialService")
    assert "processCheckout" in service_chunk.text
    assert "approved" in service_chunk.text

    # Go via tree-sitter
    go_path = os.path.join(corpus, "code/standard.go")
    go_chunks = auto_loader.load(go_path)
    assert len(go_chunks) == 3
    go_names = [c.metadata.get("name") for c in go_chunks]
    assert "CalculateMargin" in go_names
    assert "ComputeTax" in go_names

    # Check Go function body integrity
    margin_chunk = next(c for c in go_chunks if c.metadata.get("name") == "CalculateMargin")
    assert "return (revenue - cost) / revenue" in margin_chunk.text


def test_large_scale_malformed_code_fallback(corpus):
    """
    Verify that large-scale malformed code (>1MB with scattered syntax errors)
    produces reasonably-sized, non-degenerate chunks via AST fallback
    rather than failing or producing one giant fallback blob.
    """
    code_loader = CodeLoader()
    large_malformed_path = os.path.join(corpus, "code/large_malformed.py")
    assert os.path.exists(large_malformed_path)
    assert os.path.getsize(large_malformed_path) > 1_000_000

    chunks = code_loader.load(large_malformed_path)
    assert 400 <= len(chunks) <= 1200

    sizes = [len(c.text) for c in chunks]
    assert min(sizes) > 100  # No empty/degenerate chunks
    assert max(sizes) < 10_000  # No giant fallback blobs


def test_structured_loader_row_coalescing(corpus):
    """
    Verify that StructuredLoader coalesces rows into bounded chunks
    (50-100 rows per chunk) for large CSV files while keeping small files 1:1.
    """
    loader = StructuredLoader()
    # 1. Large CSV (>100k rows) coalesces into 50-row chunks
    large_csv = os.path.join(corpus, "structured/large.csv")
    large_chunks = loader.load(large_csv)
    assert 1000 <= len(large_chunks) <= 3000
    assert "Rows 1-50" in large_chunks[0].page_or_section
    assert large_chunks[0].metadata["row_count"] == 50

    # 2. Small CSV (3 rows) preserves 1 row per chunk
    small_csv = os.path.join(corpus, "data/products.csv")
    small_chunks = loader.load(small_csv)
    assert len(small_chunks) == 3


def test_async_ingestion_job_manager(corpus):
    """
    Verify IngestionJobManager threshold dispatch and job status tracking.
    """
    import tempfile
    import time
    from ico_cache.async_ingest import IngestionJobManager
    from ico_cache.core.cache_engine import CacheEngine
    from ico_cache.backends.embedding.fastembed_embedder import FastEmbedder
    from ico_cache.backends.vector.lancedb_store import LanceDBStore
    from ico_cache.backends.exact.sqlite_store import SQLiteStore

    with tempfile.TemporaryDirectory() as td:
        engine = CacheEngine(
            embedder=FastEmbedder(),
            vector_store=LanceDBStore(uri=f"{td}/lancedb"),
            exact_store=SQLiteStore(db_path=f"{td}/exact.db"),
        )
        mgr = IngestionJobManager(async_threshold_bytes=1000)

        # Sync dispatch (< 1000 bytes)
        small_file = os.path.join(corpus, "text/standard.txt")
        job_sync = mgr.submit_ingest(small_file, "tenant_test", engine)
        assert job_sync.is_async is False
        assert job_sync.status == "completed"
        assert job_sync.chunks_processed == 3

        # Async dispatch (> 1000 bytes)
        code_file = os.path.join(corpus, "code/standard.py")
        job_async = mgr.submit_ingest(code_file, "tenant_test", engine)
        assert job_async.is_async is True
        assert job_async.status in ["queued", "processing", "completed"]

        # Wait for async completion
        for _ in range(50):
            polled = mgr.get_job(job_async.job_id)
            if polled.status == "completed":
                break
            time.sleep(0.1)

        polled = mgr.get_job(job_async.job_id)
        assert polled.status == "completed"
        assert polled.chunks_processed >= 3


# NOTE: test_cross_type_adversarial_similarity_bounds was deliberately removed.
# It validated the deleted curated dataset (examples/test-corpus/queries/
# cross_type_adversarial.jsonl): 30 hand-tuned pairs whose similarity fell in a
# tight 0.8000-0.8499 band. That band is a property of the hand-curated pairs,
# not of the loaders, so it cannot (and should not) be synthesized. The
# synthetic cross-type evaluation now lives in eval_harness.py (via
# fixtures_gen.generate_cross_type_adversarial_queries) and still enforces the
# 0% false-hit regression bar at the 0.850 threshold.


def test_ingestion_job_manager_memory_bounds():
    """
    Verify IngestionJobManager bounded memory policy:
    - Evicts jobs older than job_ttl_seconds after completion
    - Enforces max_jobs capacity cap
    """
    import time
    from ico_cache.async_ingest import IngestionJobManager, IngestionJob

    # Test TTL eviction (e.g. 10 second TTL)
    mgr = IngestionJobManager(job_ttl_seconds=10, max_jobs=5)
    now = time.time()

    # Create 3 jobs: one completed 20s ago (expired), one completed 5s ago (active), one processing
    job1 = IngestionJob(
        job_id="job_old",
        tenant_id="t1",
        file_path="f1",
        file_size_bytes=100,
        status="completed",
        completed_at=now - 20.0,
    )
    job2 = IngestionJob(
        job_id="job_recent",
        tenant_id="t1",
        file_path="f2",
        file_size_bytes=100,
        status="completed",
        completed_at=now - 5.0,
    )
    job3 = IngestionJob(
        job_id="job_running",
        tenant_id="t1",
        file_path="f3",
        file_size_bytes=100,
        status="processing",
    )

    mgr.jobs["job_old"] = job1
    mgr.jobs["job_recent"] = job2
    mgr.jobs["job_running"] = job3

    # Cleanup with current time
    evicted = mgr.cleanup_expired_jobs(now=now)
    assert evicted == 1
    assert "job_old" not in mgr.jobs
    assert "job_recent" in mgr.jobs
    assert "job_running" in mgr.jobs

    # Test max_jobs capacity cap (cap is 5)
    for i in range(10):
        mgr.jobs[f"job_fill_{i}"] = IngestionJob(
            job_id=f"job_fill_{i}",
            tenant_id="t1",
            file_path=f"f_{i}",
            file_size_bytes=100,
            status="completed",
            completed_at=now - (10 - i),
        )

    mgr.cleanup_expired_jobs(now=now)
    assert len(mgr.jobs) <= 5