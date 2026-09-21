"""Office Open XML loader: .docx, .xlsx, .pptx (optional dependencies).

Each format is optional and imported lazily; a missing library degrades that
format to a no-op instead of breaking universal ingestion.
"""
import logging
import os
from typing import List, Optional

from ._common import make_chunk
from .base import BaseLoader, Chunk
from ..core.metadata_guard import MetadataSchema

logger = logging.getLogger("ico_cache.loaders.office_loader")

_MAX_CHUNK_CHARS = 2000


class OfficeLoader(BaseLoader):
    def __init__(self, schema: Optional[MetadataSchema] = None, max_chunk_chars: int = _MAX_CHUNK_CHARS):
        self.schema = schema
        self.max_chunk_chars = max_chunk_chars

    def load(self, file_path: str, schema: Optional[MetadataSchema] = None) -> List[Chunk]:
        effective_schema = schema or self.schema
        if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
            return []
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == ".docx":
                return self._load_docx(file_path, effective_schema)
            if ext == ".xlsx":
                return self._load_xlsx(file_path, effective_schema)
            if ext == ".pptx":
                return self._load_pptx(file_path, effective_schema)
        except ImportError as e:
            logger.warning("Optional office dependency missing for %s: %s", file_path, e)
        except Exception as e:
            logger.warning("Failed to parse office document %s: %s", file_path, e)
        return []

    def _load_docx(self, file_path: str, schema) -> List[Chunk]:
        import docx

        document = docx.Document(file_path)
        chunks: List[Chunk] = []
        for para in document.paragraphs:
            text = (para.text or "").strip()
            if text:
                chunks.append(
                    make_chunk(
                        text,
                        file_path,
                        f"Paragraph {len(chunks) + 1}",
                        len(chunks),
                        "docx",
                        schema,
                    )
                )
        for t_idx, table in enumerate(document.tables, start=1):
            rows = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            headers = [h or f"col{j + 1}" for j, h in enumerate(rows[0])]
            for r_idx, row in enumerate(rows[1:], start=1):
                text = "\n".join(
                    f"{headers[j] if j < len(headers) else f'col{j + 1}'}: {cell}"
                    for j, cell in enumerate(row)
                    if cell
                )
                if text:
                    chunks.append(
                        make_chunk(text, file_path, f"Table {t_idx} Row {r_idx}", len(chunks), "docx_table", schema)
                    )
        return chunks

    def _load_xlsx(self, file_path: str, schema) -> List[Chunk]:
        import openpyxl

        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        chunks: List[Chunk] = []
        try:
            for sheet in workbook.worksheets:
                rows = [
                    ["" if cell is None else str(cell) for cell in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if any(row)]
                if not rows:
                    continue
                headers = [h or f"col{j + 1}" for j, h in enumerate(rows[0])]
                for r_idx, row in enumerate(rows[1:], start=1):
                    text = "\n".join(
                        f"{headers[j] if j < len(headers) else f'col{j + 1}'}: {cell}"
                        for j, cell in enumerate(row)
                        if cell
                    )
                    if text:
                        chunks.append(
                            make_chunk(
                                text,
                                file_path,
                                f"Sheet {sheet.title} Row {r_idx}",
                                len(chunks),
                                "xlsx",
                                schema,
                            )
                        )
        finally:
            workbook.close()
        return chunks

    def _load_pptx(self, file_path: str, schema) -> List[Chunk]:
        from pptx import Presentation

        presentation = Presentation(file_path)
        chunks: List[Chunk] = []
        for s_idx, slide in enumerate(presentation.slides, start=1):
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text and text.strip():
                    chunks.append(
                        make_chunk(
                            text.strip(),
                            file_path,
                            f"Slide {s_idx}",
                            len(chunks),
                            "pptx",
                            schema,
                        )
                    )
        return chunks

    def parse(self, file_path: str, schema: Optional[MetadataSchema] = None) -> List[Chunk]:
        return self.load(file_path, schema=schema)
